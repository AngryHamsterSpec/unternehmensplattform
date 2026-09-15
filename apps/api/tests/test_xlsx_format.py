"""Synthetische XLSX-Dateien: Typen, Blattwahl, Integrität und XML-Grenzen."""

from datetime import datetime
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import Workbook
from platform_app.data.engine import DataError
from platform_app.data.schemas import ImportOptions, Step
from platform_app.data.stream_engine import build
from platform_app.data.xlsx_format import tabular_rows


def workbook(path):
    book = Workbook()
    book.active.append(["Hinweis"])
    book.active.append(["Synthetisches Beispiel"])
    sheet = book.create_sheet("Messwerte")
    sheet.append(["Name", "Wert", "Aktiv", "Zeit"])
    sheet.append([" Anna ", 1.25, True, datetime(2026, 9, 12, 8, 30)])
    sheet.append(["Bob", None, False])
    book.save(path)
    book.close()


def rewrite(path, member, transform):
    with ZipFile(path) as archive:
        parts = {name: archive.read(name) for name in archive.namelist()}
    parts[member] = transform(parts.get(member, b""))
    with ZipFile(path, "w", ZIP_DEFLATED) as archive:
        for name, value in parts.items():
            archive.writestr(name, value)


def test_sheet_types_blanks_preview_and_unchanged_original(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook(source)
    original = source.read_bytes()
    rows = list(tabular_rows(source, 2))
    assert rows[1] == [" Anna ", "1.25", "true", "2026-09-12T08:30:00"]
    assert rows[2] == ["Bob", "", "false", ""]
    normalized, _, profile, _ = build(
        source,
        tmp_path,
        ";",
        False,
        [],
        lambda *a: None,
        csv_options=ImportOptions(worksheet=2),
        source_format="xlsx",
    )
    assert profile.source_worksheet == 2 and profile.rows == 2
    preview = tmp_path / "preview"
    preview.mkdir()
    _, exported, derived, _ = build(
        normalized,
        preview,
        ";",
        True,
        [Step(operation="trim", column="Name")],
        lambda *a: None,
        source_format="xlsx",
    )
    assert derived.sample[0][0] == "Anna"
    assert b'"Anna";"1.25"' in exported.read_bytes()
    assert source.read_bytes() == original


def test_incorrect_dimensions_do_not_truncate_rows(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook(source)
    rewrite(source, "xl/worksheets/sheet2.xml", lambda b: b.replace(b'ref="A1:D3"', b'ref="A1:A1"'))
    assert len(list(tabular_rows(source, 2))) == 3


def test_headerless_keeps_first_row(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook(source)
    _, _, profile, _ = build(
        source,
        tmp_path,
        ";",
        False,
        [],
        lambda *a: None,
        csv_options=ImportOptions(worksheet=2, has_header=False),
        source_format="xlsx",
    )
    assert profile.rows == 3
    assert profile.columns[0].name == "Spalte_1"


@pytest.mark.parametrize("value,message", [("=1+2", "Formeln"), ("#DIV/0!", "Fehlerwerte")])
def test_formula_and_excel_error_are_explicit(tmp_path, value, message):
    source = tmp_path / "source.xlsx"
    book = Workbook()
    book.active.append(["Wert"])
    book.active.append([value])
    book.save(source)
    with pytest.raises(DataError, match=message):
        list(tabular_rows(source))


@pytest.mark.parametrize("sheet", [0, 3, 101])
def test_nonexistent_sheet(tmp_path, sheet):
    source = tmp_path / "source.xlsx"
    workbook(source)
    with pytest.raises(DataError, match="Arbeitsblatt"):
        list(tabular_rows(source, sheet))


def test_hidden_sheet(tmp_path):
    source = tmp_path / "source.xlsx"
    workbook(source)
    rewrite(
        source,
        "xl/workbook.xml",
        lambda b: b.replace(
            b'name="Messwerte" sheetId="2" state="visible"',
            b'name="Messwerte" sheetId="2" state="hidden"',
        ),
    )
    with pytest.raises(DataError, match="sichtbares"):
        list(tabular_rows(source, 2))


@pytest.mark.parametrize(
    "member,value,message",
    [
        ("xl/externalLinks/externalLink1.xml", b"<x/>", "externen"),
        ("../escape.xml", b"<x/>", "ZIP-Pfade"),
        ("xl/sharedStrings.xml", b"x" * (4 * 1024 * 1024 + 1), "Texttabelle"),
        ("xl/worksheets/sheet1.xml", b'<!DOCTYPE x [<!ENTITY value "bad">]><x>&value;</x>', "DTD"),
        (
            "xl/worksheets/sheet1.xml",
            b'<worksheet><sheetData><row r="1"><c r="ZZZ1"><v>1</v></c></row></sheetData></worksheet>',
            "256",
        ),
        (
            "xl/worksheets/sheet1.xml",
            b'<worksheet><sheetData><row r="999999999"><c r="A1"/></row></sheetData></worksheet>',
            "Zeilennummer",
        ),
    ],
    ids=["external-link", "zip-path", "shared-strings", "xml-dtd", "column-limit", "row-limit"],
)
def test_archive_and_xml_guards(tmp_path, member, value, message):
    source = tmp_path / "source.xlsx"
    workbook(source)
    rewrite(source, member, lambda b: value)
    with pytest.raises(DataError, match=message):
        list(tabular_rows(source))


def test_invalid_archive_and_cancellation(tmp_path):
    source = tmp_path / "source.xlsx"
    source.write_bytes(b"no zip")
    with pytest.raises(DataError, match="beschädigt"):
        list(tabular_rows(source))
    workbook(source)

    def cancel():
        raise RuntimeError("Testabbruch")

    with pytest.raises(RuntimeError, match="Testabbruch"):
        list(tabular_rows(source, heartbeat=cancel))
