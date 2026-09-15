"""Begrenzter XLSX-Adapter: ZIP prüfen, Werte lesen, keine Formel-/Linkausführung."""

import math
import re
import struct
from collections.abc import Callable, Iterator
from datetime import date, datetime, time
from pathlib import Path, PurePosixPath
from xml.etree.ElementTree import ParseError
from xml.parsers.expat import ExpatError, ParserCreate
from zipfile import BadZipFile, ZipFile

from defusedxml.common import DefusedXmlException  # type: ignore[import-untyped]
from openpyxl import load_workbook  # type: ignore[import-untyped]
from openpyxl.utils.exceptions import InvalidFileException  # type: ignore[import-untyped]

from platform_app.data.engine import DataError

MAX_EXPANDED = 2 * 1024 * 1024 * 1024
MAX_METADATA = 16 * 1024 * 1024
MAX_STRINGS = 4 * 1024 * 1024


def inspect_xml(archive: ZipFile, name: str, heartbeat: Callable[[], None]) -> None:
    # SAX-Vorprüfung begrenzt auch einzelne XML-Tokens, bevor openpyxl Objekte anlegt.
    parser = ParserCreate(namespace_separator="}")
    depth = 0
    cell_chars = 0
    row_chars = 0
    in_cell = False
    nodes = 0
    row_number = 0
    previous_column = 0
    sheet = name.startswith("xl/worksheets/") and name.endswith(".xml")

    def forbidden(*args: object) -> None:
        raise DataError("XLSX mit DTD oder XML-Entitäten wird nicht unterstützt.")

    def start(tag: str, attributes: dict[str, str]) -> None:
        nonlocal depth, cell_chars, row_chars, in_cell, nodes, row_number, previous_column
        depth += 1
        nodes += 1
        if depth > 64 or (not sheet and nodes > 100000):
            raise DataError("Die XLSX-XML-Struktur ist zu komplex.")
        tag = tag.rsplit("}", 1)[-1]
        if sheet and tag == "row":
            row_chars = 0
            previous_column = 0
            new_row = int(attributes.get("r", str(row_number + 1)))
            if not row_number < new_row <= 1048576:
                raise DataError(
                    "Eine XLSX-Zeilennummer ist doppelt, unsortiert oder liegt außerhalb des Excel-Formats."
                )
            row_number = new_row
        if sheet and tag == "c":
            cell_chars, in_cell = 0, True
            ref = attributes.get("r", "")
            match = re.fullmatch(r"([A-Z]{1,3})([1-9][0-9]{0,6})", ref)
            if ref and not match:
                raise DataError("Eine XLSX-Zelladresse ist ungültig oder fehlt.")
            column = previous_column + 1
            if match:
                column = 0
                for char in match[1]:
                    column = column * 26 + ord(char) - 64
                if int(match[2]) != row_number or column <= previous_column:
                    raise DataError(
                        "XLSX-Zelladressen sind doppelt, unsortiert oder passen nicht zur Zeile."
                    )
            previous_column = column
            if column > 256:
                raise DataError(
                    "Das Arbeitsblatt überschreitet 256 Spalten oder die Excel-Zeilengrenze. Bitte den Datenbereich als CSV exportieren."
                )

    def end(tag: str) -> None:
        nonlocal depth, in_cell
        depth -= 1
        if tag.rsplit("}", 1)[-1] == "c":
            in_cell = False

    def characters(value: str) -> None:
        nonlocal cell_chars, row_chars
        if in_cell:
            cell_chars += len(value)
            row_chars += len(value.encode("utf-8"))
            if cell_chars > 1048576 or row_chars > 4 * 1024 * 1024:
                raise DataError(
                    "Eine XLSX-Zelle überschreitet 1.048.576 Zeichen oder eine Zeile 4 MiB."
                )

    parser.StartElementHandler = start
    parser.EndElementHandler = end
    parser.CharacterDataHandler = characters
    parser.StartDoctypeDeclHandler = forbidden
    parser.EntityDeclHandler = forbidden
    with archive.open(name) as stream:
        consumed = 0
        while chunk := stream.read(65536):
            consumed += len(chunk)
            parser.Parse(chunk, False)
            if consumed - parser.CurrentByteIndex > 4 * 1024 * 1024:
                raise DataError("Ein XLSX-XML-Element überschreitet 4 MiB.")
            heartbeat()
        parser.Parse(b"", True)


def inspect_archive(path: Path, heartbeat: Callable[[], None]) -> None:
    # ZIP lädt das Zentralverzeichnis in den RAM: vor dessen Konstruktion begrenzen.
    with path.open("rb") as source:
        size = source.seek(0, 2)
        source.seek(max(0, size - 65557))
        tail = source.read(65557)
    offset = tail.rfind(b"PK\x05\x06")
    if offset < 0 or offset + 22 > len(tail):
        raise DataError("Die XLSX-ZIP-Struktur ist beschädigt.")
    _, disk, central_disk, disk_entries, entries, central_size, _, comment_size = (
        struct.unpack_from("<4s4H2LH", tail, offset)
    )
    if (
        disk
        or central_disk
        or disk_entries != entries
        or entries > 2048
        or central_size > 2 * 1024 * 1024
        or offset + 22 + comment_size != len(tail)
    ):
        raise DataError(
            "Das XLSX-Zentralverzeichnis ist ungültig oder zu groß (2.048 Einträge / 2 MiB). ZIP64 wird nicht unterstützt."
        )
    with ZipFile(path) as archive:
        members = archive.infolist()
        if len(members) > 2048 or sum(m.file_size for m in members) > MAX_EXPANDED:
            raise DataError(
                "Die XLSX-Datei überschreitet die Grenze von 2 GiB entpackt oder 2.048 ZIP-Einträgen."
            )
        seen = set()
        metadata = 0
        for member in members:
            name = member.filename
            lower = name.lower()
            parts = PurePosixPath(name)
            if name in seen or parts.is_absolute() or ".." in parts.parts or "\\" in name:
                raise DataError("Die XLSX-Datei enthält doppelte oder unzulässige ZIP-Pfade.")
            seen.add(name)
            if member.flag_bits & 1:
                raise DataError("Verschlüsselte XLSX-Dateien werden nicht unterstützt.")
            if any(
                part in lower
                for part in ("vbaproject", "externallinks/", "embeddings/", "activex/")
            ):
                raise DataError(
                    "XLSX mit Makros, eingebetteten Objekten oder externen Verknüpfungen wird nicht importiert."
                )
            if lower == "xl/sharedstrings.xml" and member.file_size > MAX_STRINGS:
                raise DataError(
                    "Die gemeinsame XLSX-Texttabelle überschreitet 4 MiB. Bitte als CSV exportieren."
                )
            if member.file_size > max(1024 * 1024, member.compress_size * 1000):
                raise DataError("Die XLSX-Datei hat ein unzulässig hohes Entpackverhältnis.")
            if not (lower.startswith("xl/worksheets/") and lower.endswith(".xml")):
                metadata += member.file_size
        if metadata > MAX_METADATA:
            raise DataError("XLSX-Metadaten überschreiten 16 MiB. Bitte als CSV exportieren.")
        for member in members:
            if member.filename.endswith((".xml", ".rels")):
                inspect_xml(archive, member.filename, heartbeat)
        if "[Content_Types].xml" not in seen or "xl/workbook.xml" not in seen:
            raise DataError("Die Datei ist kein unterstütztes XLSX-Arbeitsbuch.")


def cell_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        raise DataError("Eine Excel-Zelle enthält keine endliche Zahl.")
    if isinstance(value, (str, int, float)):
        return str(value)
    raise DataError("Ein Excel-Zelltyp wird nicht als Tabellenwert unterstützt.")


def tabular_rows(
    path: Path, worksheet: int = 1, heartbeat: Callable[[], None] = lambda: None
) -> Iterator[list[str]]:
    book = None
    try:
        inspect_archive(path, heartbeat)
        # Keine Cachewerte von Formeln, externe Referenzen oder Makros übernehmen.
        with path.open("rb") as source:
            book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
            if not 1 <= worksheet <= min(100, len(book.worksheets)):
                raise DataError(
                    f"Arbeitsblatt {worksheet} fehlt. Die Datei enthält {len(book.worksheets)} Arbeitsblätter."
                )
            sheet = book.worksheets[worksheet - 1]
            if sheet.sheet_state != "visible":
                raise DataError("Bitte ein sichtbares Arbeitsblatt auswählen.")
            sheet.reset_dimensions()
            width = None
            for row in sheet.iter_rows():
                heartbeat()
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        raise DataError(
                            "Das Arbeitsblatt enthält Formeln. Bitte die gewünschten Ergebnisse in Excel als Werte speichern; Formel-Cachewerte werden nicht ungeprüft übernommen."
                        )
                    if cell.data_type == "e":
                        raise DataError(
                            "Das Arbeitsblatt enthält Excel-Fehlerwerte. Bitte diese vor dem Import korrigieren."
                        )
                    values.append(cell_text(cell.value))
                # Formatierte leere Randzellen definieren keine zusätzlichen Spalten.
                while values and values[-1] == "":
                    values.pop()
                if width is None:
                    if not values:
                        continue
                    width = len(values)
                    if width > 256:
                        raise DataError("Das Arbeitsblatt überschreitet 256 Datenspalten.")
                if len(values) > width:
                    raise DataError(
                        "Eine Excel-Zeile enthält zusätzliche Datenspalten gegenüber der ersten Tabellenzeile."
                    )
                yield values + [""] * (width - len(values))
    except DataError:
        raise
    except (
        BadZipFile,
        InvalidFileException,
        ParseError,
        ExpatError,
        DefusedXmlException,
        KeyError,
        ValueError,
        OSError,
        TypeError,
        IndexError,
    ):
        raise DataError(
            "Die XLSX-Struktur ist beschädigt oder wird nicht unterstützt. Bitte als normale XLSX- oder CSV-Datei neu speichern."
        ) from None
    finally:
        if book is not None:
            book.close()
